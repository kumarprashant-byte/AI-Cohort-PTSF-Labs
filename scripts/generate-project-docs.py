"""Convert delivery/*.md project docs to .docx."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
import re

SOURCES = [
    ("delivery/project-charter.md", "delivery/PTSF-Project-Charter.docx"),
    ("delivery/raid-log.md", "delivery/PTSF-RAID-Log.docx"),
    ("delivery/raci-matrix.md", "delivery/PTSF-RACI-Matrix.docx"),
    ("delivery/status-report.md", "delivery/PTSF-Status-Report.docx"),
]

def strip_inline(text):
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'`(.+?)`', r'\1', text)
    text = re.sub(r'\*(.+?)\*', r'\1', text)
    return text

def convert(md_path, out_path):
    lines = Path(md_path).read_text(encoding='utf-8').splitlines()
    doc = Document()
    in_front = False
    front_done = False
    i = 0
    while i < len(lines):
        line = lines[i]
        if i == 0 and line.strip() == '---':
            in_front = True; i += 1; continue
        if in_front and line.strip() == '---':
            in_front = False; front_done = True; i += 1; continue
        if in_front:
            i += 1; continue

        if line.startswith('# '):
            doc.add_heading(strip_inline(line[2:].strip()), level=0)
        elif line.startswith('## '):
            doc.add_heading(strip_inline(line[3:].strip()), level=1)
        elif line.startswith('### '):
            doc.add_heading(strip_inline(line[4:].strip()), level=2)
        elif line.startswith('|'):
            # collect table block
            tbl = []
            while i < len(lines) and lines[i].startswith('|'):
                tbl.append(lines[i]); i += 1
            # parse
            rows = []
            for r in tbl:
                cells = [c.strip() for c in r.strip().strip('|').split('|')]
                if all(re.fullmatch(r':?-+:?', c) for c in cells):
                    continue
                rows.append(cells)
            if rows:
                t = doc.add_table(rows=len(rows), cols=len(rows[0]))
                t.style = 'Light Grid Accent 1'
                for r_idx, row in enumerate(rows):
                    for c_idx, val in enumerate(row):
                        if c_idx < len(t.rows[r_idx].cells):
                            t.rows[r_idx].cells[c_idx].text = strip_inline(val)
            continue
        elif line.startswith('- ') or line.startswith('* '):
            doc.add_paragraph(strip_inline(line[2:].strip()), style='List Bullet')
        elif re.match(r'^\d+\.\s', line):
            doc.add_paragraph(strip_inline(re.sub(r'^\d+\.\s', '', line)), style='List Number')
        elif line.strip() == '':
            pass
        else:
            doc.add_paragraph(strip_inline(line))
        i += 1
    doc.save(out_path)
    print(f"Wrote {out_path}")

for src, dst in SOURCES:
    convert(src, dst)
