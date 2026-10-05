"""Generate PTSF-Manual-Tests.docx from the pending manual criteria."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# ---- Title ----
title = doc.add_heading("PTSF Patient Travel Support — Manual Test Checklist", level=0)
p = doc.add_paragraph()
p.add_run("Target org: ").bold = True
p.add_run("epic.out.e68d765f4115@orgfarm.salesforce.com (alias ptsf-lab)\n")
p.add_run("Generated: ").bold = True
p.add_run("2026-09-30\n")
p.add_run("Companion doc: ").bold = True
p.add_run("delivery/test-plan.md (full walkthrough with CLI commands and record ids)")

doc.add_paragraph(
    "30 manual sign-offs are pending across 15 delivered intents. "
    "Each item names the intent and criterion (join key for sign-off). "
    "Record results back with /ql-record-test-execution INT-NNN or by editing "
    "the Sign-off cell in intents/INT-NNN/test-script.md."
)

# ---- Section A: doable now ----
doc.add_heading("A. Doable now in the sandbox (System Admin)", level=1)
doc.add_paragraph(
    "15 criteria across 8 intents. All prerequisites satisfied: metadata deployed, "
    "seed data loaded. Full step-by-step commands in delivery/test-plan.md §3.2–§3.6."
)

section_a = [
    ("INT-006", "C8", "Localized welcome email fires after Summary step of Patient Portal wizard.",
     "Complete the wizard as an EMEA patient (French) → confirm email received in French. Check Setup → Email Logs if inbox not available."),
    ("INT-006", "C9", "Wizard entry publishes HIC prefill event without blocking UI (guardrail 1).",
     "Open wizard step 1 → confirm HIC_Prefill_Requested__e appears in Setup → Platform Events → Recent Events, and the wizard UI does not freeze while it publishes."),
    ("INT-008", "C7", "Rejection email is sent to the patient on the Covered branch.",
     "After Scene A of C12 completes, check Setup → Email Logs — one outbound email to the patient Contact, subject 'Your subsidy application'."),
    ("INT-008", "C12", "Auto-reject when HIC returns Covered (end-to-end).",
     "1) Open AMER Submitted app (a05oB000002vrG9QAI). 2) Change Status to Submitted, save → confirm advance to Awaiting_Insurance_Check within 5s, Submitted_At__c stamped, Medical_History_Snapshot__c populated. 3) Publish HIC_Entitlement_Completed__e with Source_Record_Id__c=<app id> Status__c=Covered. 4) Confirm Status=Rejected, Rejection_Reason__c populated, email queued."),
    ("INT-008", "C13", "Not-covered branch routes to Awaiting_Practitioner.",
     "Same as C12 but publish HIC_Entitlement_Completed__e with Status__c=Not_Covered → confirm Status=Awaiting_Practitioner, no rejection email."),
    ("INT-010", "C8", "SLA reassignment fires after 3 business days.",
     "1) Take APAC assignment (a03oB000000eidRQAQ). 2) Backdate Assigned_Date__c to 4 days ago via `sf data update record`. 3) Setup → Flows → Assignment_SLA_Check → Run. 4) Confirm Assignment.Status=Reassigned, Application.Previously_Declined_Practitioners__c captures the practitioner Contact Id, practitioner email queued."),
    ("INT-011", "C8", "Assessment save-draft + submit round-trip, isolated per Assignment.",
     "1) Open an accepted oncology Assignment. 2) Fill Assessment LWC fields, click Save Draft → confirm one Assessment__c row Status=Draft. 3) Edit, Save Draft again → same row updated (no dup). 4) Click Submit → Status=Pending Review, Submitted_Date__c stamped. 5) Open a second Assignment for the same practitioner → form is independent, no data leakage."),
    ("INT-012", "C5", "Assessment escalation after 15 days creates a Task (once).",
     "1) Create an Assessment 16 days old (`Test.setCreatedDate` won't work in the live org; use CreatedDate override via a Batch or wait — for the manual test, use an existing aged Assessment or the seeded one). 2) Setup → Flows → Assessment_Escalation_Check → Run. 3) Confirm one Task created WhatId=Assessment, Assessment.Escalated__c=true, Escalated_Date__c=today."),
    ("INT-012", "C6", "Double-escalate guardrail: re-run creates no second Task.",
     "Re-run the Flow from C5. Confirm the Assessment still has exactly one associated Task."),
    ("INT-014", "C5", "Utility-bar item on Application_Review launches Determine_Subsidy Flow.",
     "Open Application_Review app → click the utility bar item → confirm the Determine_Subsidy Flow launches inline (Setup → App Manager → Application_Review → confirm utility item is configured)."),
    ("INT-014", "C7", "Under-threshold subsidy auto-approves.",
     "Run Determine_Subsidy flow on an EMEA app with Transport=Air, Distance=300, multiplier 1.0, threshold 500 → confirm Proposed_Amount__c=300, Approved_Amount__c=300."),
    ("INT-014", "C8", "Over-threshold subsidy requires manual approval.",
     "Same flow with Distance=1800, threshold 1500 → confirm Proposed_Amount__c=1800, Approved_Amount__c stays null, over-threshold screen note displayed."),
    ("INT-016", "C6", "No report or dashboard references Medical_History__c.",
     "Walk every report in PTSF_Reports and every dashboard component in PTSF_Dashboards — confirm none reference Medical_History__c or any medical-history field. Any reference is a P1 defect (privacy)."),
    ("INT-017", "C7", "Migration reconciliation reports show 100% count parity.",
     "Run the migration → open Migration_Reconciliation report → confirm source-vs-target counts match 100% per region."),
    ("INT-017", "C8", "In-flight applications are not migrated.",
     "Confirm the migration filter excluded applications with Status ∈ {Submitted, Awaiting_*, In_Progress} — count matches 0 in the target for those statuses in the source."),
]

for intent, crit, summary, steps in section_a:
    p = doc.add_paragraph()
    p.add_run(f"☐ {intent}-{crit} — ").bold = True
    p.add_run(summary)
    doc.add_paragraph(f"How: {steps}").paragraph_format.left_indent = Inches(0.3)
    doc.add_paragraph(f"Sign-off: __________________  Date: __________  Pass / Fail").paragraph_format.left_indent = Inches(0.3)

# ---- Section B: blocked pending persona provisioning ----
doc.add_heading("B. Blocked pending persona provisioning (needs SAML JIT + test users)", level=1)
doc.add_paragraph(
    "10 criteria across 6 intents. Prerequisite: PTSF IT provides EMEA/AMER/APAC AD IdP metadata + certificates, "
    "and INT-002 SAML JIT handler is configured (see runbook intents/INT-002/). Then INT-005 role + permission-set groups "
    "assign correctly via JIT."
)

section_b = [
    ("INT-002", "C6", "EMEA user SAML-authenticates and lands with Region__c=EMEA, Language=fr."),
    ("INT-005", "C6", "APAC Assessor user has correct Role, PSG, Region__c after JIT provisioning."),
    ("INT-005", "C7", "EMEA Regional Ops Manager user has correct Role, PSG, Region__c after JIT."),
    ("INT-013", "C10", "Assessor console shows My Queue, Assessment + HIC_Check tiles."),
    ("INT-015", "C5", "Einstein bot never surfaces medical-history references."),
    ("INT-015", "C6", "Patient chat routes in-region unless no in-region agent."),
    ("INT-015", "C7", "French EMEA patient bot→handoff walkthrough."),
    ("INT-016", "C7", "APAC Regional Ops Manager Team Manager Dashboard shows APAC rows only."),
    ("INT-020", "C6", "APAC Assessor SOQL returns APAC rows only; EMEA Regional Ops Manager returns EMEA only."),
    ("INT-020", "C7", "Same personas open All Open Applications list view → same-region rows only."),
]

for intent, crit, summary in section_b:
    p = doc.add_paragraph()
    p.add_run(f"☐ {intent}-{crit} — ").bold = True
    p.add_run(summary)
    doc.add_paragraph(f"Sign-off: __________________  Date: __________  Pass / Fail").paragraph_format.left_indent = Inches(0.3)

# ---- Section C: external prerequisite ----
doc.add_heading("C. Blocked pending Facebook Dev App", level=1)
doc.add_paragraph("1 criterion. Prerequisite: Facebook Dev App + Auth Provider configured (runbook intents/INT-003/).")
p = doc.add_paragraph()
p.add_run("☐ INT-003-C4 — ").bold = True
p.add_run("Patient registers via Facebook → onboarding entry; Practitioner self-registers → admin approves → logs in.")
doc.add_paragraph("Sign-off: __________________  Date: __________  Pass / Fail").paragraph_format.left_indent = Inches(0.3)

# ---- Section D: cutover-time ----
doc.add_heading("D. Production cutover-time observations", level=1)
doc.add_paragraph(
    "4 criteria across 1 intent (INT-018). These are cutover-weekend observations, "
    "not a dev-cycle test. Run during the actual TAMS migration."
)

section_d = [
    ("INT-018", "C2", "Freeze window enforced for TAMS new-application intake."),
    ("INT-018", "C3", "4-week dual-run window observed; no early closure."),
    ("INT-018", "C4", "Only one region cut over per weekend."),
    ("INT-018", "C6", "AMER pilot cutover acceptance walkthrough."),
]
for intent, crit, summary in section_d:
    p = doc.add_paragraph()
    p.add_run(f"☐ {intent}-{crit} — ").bold = True
    p.add_run(summary)
    doc.add_paragraph(f"Sign-off: __________________  Date: __________  Pass / Fail").paragraph_format.left_indent = Inches(0.3)

# ---- Footer ----
doc.add_heading("Recording sign-offs", level=1)
doc.add_paragraph(
    "For each pass: run  /ql-record-test-execution INT-NNN  — writes a persistent execution report under "
    "intents/INT-NNN/test-evidence/ and fills the Sign-off cell in intents/INT-NNN/test-script.md."
)
doc.add_paragraph(
    "For each fail: capture actual vs expected, then run  /ql-diagnose  to route it to a fix branch."
)

out = "delivery/PTSF-Manual-Tests.docx"
doc.save(out)
print(f"Wrote {out}")
